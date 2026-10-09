"""
StudyMate AI - FastAPI Backend
Complete REST API for the AI tutoring platform
"""
import os
import uuid
import json
import asyncio
from contextlib import asynccontextmanager
from datetime import datetime, timezone, timedelta
from typing import List, Optional

from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Form, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from dotenv import load_dotenv

load_dotenv()

from app.database import (
    init_db, get_db, User, Course, Document, DocumentChunk,
    Topic, Concept, TopicPrerequisite, Assessment, Question,
    QuestionAttempt, Mastery, TutorSession, TutorMessage,
    StudyPlan, EvaluationResult
)
from app.seed_data import (
    DEMO_USER_ID, DEMO_COURSE_ID,
    OS_TEXTBOOK_CHUNKS, OS_SLIDES_CHUNKS, OS_VIDEO_CHUNKS,
    get_demo_documents, get_demo_topics, get_topic_prerequisites,
    get_demo_mastery, EVALUATION_QUESTIONS
)
from app.services.rag_service import rag_service
from app.services.assessment_service import assessment_service
from app.services.learner_model import learner_model_service
from app.services.evaluation_service import evaluation_service

@asynccontextmanager
async def lifespan(app: FastAPI):
    """Startup/shutdown lifespan handler."""
    init_db()
    db = next(get_db())
    try:
        await seed_demo_data(db)
    finally:
        db.close()
    yield  # application runs here

app = FastAPI(title="StudyMate AI", version="1.0.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ============================================================
# Pydantic Models
# ============================================================

class LoginRequest(BaseModel):
    mode: str = "demo"  # "demo" or "credentials"
    email: Optional[str] = None
    password: Optional[str] = None

class CourseCreate(BaseModel):
    name: str
    description: str = ""

class TutorQuery(BaseModel):
    course_id: str
    question: str
    session_id: Optional[str] = None
    language: str = "en"

class AssessmentCreate(BaseModel):
    course_id: str
    topics: List[str]
    difficulty: str = "medium"
    question_count: int = 5
    question_types: List[str] = ["mcq", "short_answer"]

class AnswerSubmit(BaseModel):
    answers: dict  # question_id -> user_answer

class MasteryUpdate(BaseModel):
    course_id: str
    topic: str
    is_correct: bool
    difficulty: str = "medium"

class StudyPlanCreate(BaseModel):
    course_id: str
    exam_date: Optional[str] = None
    daily_hours: float = 2.0

class RevisionRequest(BaseModel):
    course_id: str
    topic: str
    type: str  # flashcards, notes, summary

# ============================================================
# Demo data seeding (called from lifespan)
# ============================================================


async def seed_demo_data(db: Session):
    """Seed the database with demo data if not already present."""
    existing = db.query(User).filter(User.id == DEMO_USER_ID).first()
    if existing:
        # If DB already has data but in-memory RAG service is empty, reload chunks
        if not rag_service.chunks:
            chunks = db.query(DocumentChunk).all()
            for c in chunks:
                emb = await rag_service.get_embedding(c.content)
                rag_service.index_chunk(c.id, {
                    "course_id": c.course_id,
                    "document_id": c.document_id,
                    "content": c.content,
                    "source_type": c.source_type,
                    "source_name": c.source_name,
                    "page_number": c.page_number,
                    "slide_number": c.slide_number,
                    "timestamp": c.timestamp,
                    "topic": c.topic,
                    "subtopic": c.subtopic,
                    "concept": c.concept,
                    "difficulty": c.difficulty
                }, emb)
            print(f"[OK] Reloaded {len(rag_service.chunks)} existing chunks into RAG service")
        return

    # Create demo user
    user = User(
        id=DEMO_USER_ID,
        name="Thanusha",
        email="thanusha@studymate.demo",
        language="en"
    )
    db.add(user)

    # Create demo course
    course = Course(
        id=DEMO_COURSE_ID,
        user_id=DEMO_USER_ID,
        name="Operating Systems",
        description="Semester 5 Operating Systems — Covers processes, threads, synchronization, deadlocks, and memory management.",
        is_demo=True
    )
    db.add(course)

    # Create documents
    for doc_data in get_demo_documents():
        doc = Document(
            id=doc_data["id"],
            course_id=DEMO_COURSE_ID,
            name=doc_data["name"],
            file_type=doc_data["file_type"],
            file_path=f"/demo/{doc_data['name']}",
            file_size=doc_data["file_size"],
            status="completed",
            processing_progress=100.0
        )
        db.add(doc)

    # Create topics and concepts
    for topic_data in get_demo_topics():
        topic = Topic(
            id=topic_data["id"],
            course_id=DEMO_COURSE_ID,
            name=topic_data["name"],
            unit=topic_data["unit"],
            description=topic_data["description"],
            difficulty=topic_data["difficulty"],
            order_index=topic_data["order_index"]
        )
        db.add(topic)

        # Create subtopics
        for st in topic_data.get("subtopics", []):
            sub = Topic(
                id=st["id"],
                course_id=DEMO_COURSE_ID,
                name=st["name"],
                unit=topic_data["unit"],
                parent_topic_id=topic_data["id"],
                order_index=0
            )
            db.add(sub)

        # Create concepts
        for concept_name in topic_data.get("concepts", []):
            concept = Concept(
                id=f"concept-{uuid.uuid4().hex[:8]}",
                topic_id=topic_data["id"],
                name=concept_name
            )
            db.add(concept)

    # Create prerequisites
    for topic_id, prereq_id in get_topic_prerequisites():
        prereq = TopicPrerequisite(
            id=f"prereq-{uuid.uuid4().hex[:8]}",
            topic_id=topic_id,
            prerequisite_topic_id=prereq_id
        )
        db.add(prereq)

    # Create chunks and index them for RAG
    chunk_index = 0
    # Textbook chunks
    for chunk_data in OS_TEXTBOOK_CHUNKS:
        chunk_id = f"chunk-{uuid.uuid4().hex[:8]}"
        chunk = DocumentChunk(
            id=chunk_id,
            document_id="doc-textbook-001",
            course_id=DEMO_COURSE_ID,
            content=chunk_data["content"],
            source_type="pdf",
            source_name="Operating Systems Textbook",
            page_number=chunk_data["page_number"],
            topic=chunk_data["topic"],
            subtopic=chunk_data["subtopic"],
            concept=chunk_data["concept"],
            difficulty=chunk_data["difficulty"],
            chunk_index=chunk_index
        )
        db.add(chunk)

        # Index for RAG
        embedding = await rag_service.get_embedding(chunk_data["content"])
        rag_service.index_chunk(chunk_id, {
            "course_id": DEMO_COURSE_ID,
            "document_id": "doc-textbook-001",
            "content": chunk_data["content"],
            "source_type": "pdf",
            "source_name": "Operating Systems Textbook",
            "page_number": chunk_data["page_number"],
            "topic": chunk_data["topic"],
            "subtopic": chunk_data["subtopic"],
            "concept": chunk_data["concept"],
            "difficulty": chunk_data["difficulty"]
        }, embedding)
        chunk_index += 1

    # Slides chunks
    for chunk_data in OS_SLIDES_CHUNKS:
        chunk_id = f"chunk-{uuid.uuid4().hex[:8]}"
        chunk = DocumentChunk(
            id=chunk_id,
            document_id="doc-slides-001",
            course_id=DEMO_COURSE_ID,
            content=chunk_data["content"],
            source_type="pptx",
            source_name="OS Lecture Slides",
            slide_number=chunk_data["slide_number"],
            topic=chunk_data["topic"],
            subtopic=chunk_data["subtopic"],
            concept=chunk_data["concept"],
            difficulty=chunk_data["difficulty"],
            chunk_index=chunk_index
        )
        db.add(chunk)

        embedding = await rag_service.get_embedding(chunk_data["content"])
        rag_service.index_chunk(chunk_id, {
            "course_id": DEMO_COURSE_ID,
            "document_id": "doc-slides-001",
            "content": chunk_data["content"],
            "source_type": "pptx",
            "source_name": "OS Lecture Slides",
            "slide_number": chunk_data["slide_number"],
            "topic": chunk_data["topic"],
            "subtopic": chunk_data["subtopic"],
            "concept": chunk_data["concept"],
            "difficulty": chunk_data["difficulty"]
        }, embedding)
        chunk_index += 1

    # Video chunks
    for chunk_data in OS_VIDEO_CHUNKS:
        chunk_id = f"chunk-{uuid.uuid4().hex[:8]}"
        chunk = DocumentChunk(
            id=chunk_id,
            document_id="doc-video-001",
            course_id=DEMO_COURSE_ID,
            content=chunk_data["content"],
            source_type="video",
            source_name="OS Lecture Video",
            video_timestamp_start=chunk_data["timestamp_start"],
            video_timestamp_end=chunk_data["timestamp_end"],
            topic=chunk_data["topic"],
            subtopic=chunk_data["subtopic"],
            concept=chunk_data["concept"],
            difficulty=chunk_data["difficulty"],
            chunk_index=chunk_index
        )
        db.add(chunk)

        embedding = await rag_service.get_embedding(chunk_data["content"])
        rag_service.index_chunk(chunk_id, {
            "course_id": DEMO_COURSE_ID,
            "document_id": "doc-video-001",
            "content": chunk_data["content"],
            "source_type": "video",
            "source_name": "OS Lecture Video",
            "video_timestamp_start": chunk_data["timestamp_start"],
            "video_timestamp_end": chunk_data["timestamp_end"],
            "topic": chunk_data["topic"],
            "subtopic": chunk_data["subtopic"],
            "concept": chunk_data["concept"],
            "difficulty": chunk_data["difficulty"]
        }, embedding)
        chunk_index += 1

    # Create initial mastery
    for m in get_demo_mastery():
        mastery = Mastery(
            id=f"mastery-{uuid.uuid4().hex[:8]}",
            user_id=DEMO_USER_ID,
            course_id=DEMO_COURSE_ID,
            topic=m["topic"],
            mastery_score=m["mastery_score"],
            confidence=m["confidence"],
            attempts=m["attempts"],
            correct=m["correct"],
            last_reviewed=datetime.now(timezone.utc) - timedelta(days=random.randint(1, 7)),
            difficulty_history=[]
        )
        db.add(mastery)

    # Create a sample completed assessment
    assessment = Assessment(
        id="assessment-demo-001",
        user_id=DEMO_USER_ID,
        course_id=DEMO_COURSE_ID,
        title="Deadlocks & Synchronization Quiz",
        topics=["Deadlocks", "Synchronization"],
        difficulty="medium",
        question_count=5,
        score=72.0,
        completed=True,
        created_at=datetime.now(timezone.utc) - timedelta(days=2),
        completed_at=datetime.now(timezone.utc) - timedelta(days=2)
    )
    db.add(assessment)

    assessment2 = Assessment(
        id="assessment-demo-002",
        user_id=DEMO_USER_ID,
        course_id=DEMO_COURSE_ID,
        title="Process Management Quiz",
        topics=["Process Management"],
        difficulty="easy",
        question_count=5,
        score=88.0,
        completed=True,
        created_at=datetime.now(timezone.utc) - timedelta(days=5),
        completed_at=datetime.now(timezone.utc) - timedelta(days=5)
    )
    db.add(assessment2)

    db.commit()
    print("[OK] Demo data seeded successfully")


# ============================================================
# AUTH
# ============================================================

@app.post("/api/auth/login")
async def login(request: LoginRequest):
    """Demo login — always succeeds."""
    return {
        "user": {
            "id": DEMO_USER_ID,
            "name": "Thanusha",
            "email": "thanusha@studymate.demo",
            "avatar_url": None,
            "language": "en"
        },
        "token": "demo-token-studymate"
    }


# ============================================================
# COURSES
# ============================================================

@app.get("/api/courses")
async def get_courses(user_id: str = DEMO_USER_ID, db: Session = Depends(get_db)):
    courses = db.query(Course).filter(Course.user_id == user_id).all()
    return [{
        "id": c.id,
        "name": c.name,
        "description": c.description,
        "is_demo": c.is_demo,
        "created_at": c.created_at.isoformat() if c.created_at else None,
        "document_count": db.query(Document).filter(Document.course_id == c.id).count(),
        "topic_count": db.query(Topic).filter(Topic.course_id == c.id, Topic.parent_topic_id == None).count()
    } for c in courses]


@app.post("/api/courses")
async def create_course(data: CourseCreate, db: Session = Depends(get_db)):
    course = Course(
        id=f"course-{uuid.uuid4().hex[:8]}",
        user_id=DEMO_USER_ID,
        name=data.name,
        description=data.description
    )
    db.add(course)
    db.commit()
    return {"id": course.id, "name": course.name, "description": course.description}


@app.get("/api/courses/{course_id}")
async def get_course(course_id: str, db: Session = Depends(get_db)):
    course = db.query(Course).filter(Course.id == course_id).first()
    if not course:
        raise HTTPException(404, "Course not found")
    return {
        "id": course.id,
        "name": course.name,
        "description": course.description,
        "is_demo": course.is_demo,
        "documents": [{
            "id": d.id,
            "name": d.name,
            "file_type": d.file_type,
            "status": d.status,
            "file_size": d.file_size,
            "processing_progress": d.processing_progress
        } for d in course.documents]
    }


# ============================================================
# DOCUMENTS
# ============================================================

@app.get("/api/documents/{course_id}")
async def get_documents(course_id: str, db: Session = Depends(get_db)):
    docs = db.query(Document).filter(Document.course_id == course_id).all()
    return [{
        "id": d.id,
        "name": d.name,
        "file_type": d.file_type,
        "status": d.status,
        "file_size": d.file_size,
        "processing_progress": d.processing_progress,
        "created_at": d.created_at.isoformat() if d.created_at else None
    } for d in docs]


@app.post("/api/documents/upload")
async def upload_document(
    course_id: str = Form(...),
    file: UploadFile = File(...)
):
    """Upload and process a document."""
    ext = file.filename.split(".")[-1].lower() if file.filename else ""
    if ext not in ["pdf", "ppt", "pptx", "mp4"]:
        raise HTTPException(400, f"Unsupported file type: {ext}")

    doc_id = f"doc-{uuid.uuid4().hex[:8]}"
    # Save file
    upload_dir = os.path.join(os.path.dirname(__file__), "..", "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    file_path = os.path.join(upload_dir, f"{doc_id}.{ext}")

    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    db = next(get_db())
    try:
        doc = Document(
            id=doc_id,
            course_id=course_id,
            name=file.filename,
            file_type=ext,
            file_path=file_path,
            file_size=len(content),
            status="processing",
            processing_progress=10.0
        )
        db.add(doc)
        db.commit()

        # Process in background
        asyncio.create_task(process_document(doc_id, course_id, file_path, ext, content))

        return {
            "id": doc_id,
            "name": file.filename,
            "status": "processing",
            "file_type": ext
        }
    finally:
        db.close()


async def process_document(doc_id: str, course_id: str, file_path: str, ext: str, content: bytes):
    """Process uploaded document - extract text, create chunks, generate embeddings."""
    db = next(get_db())
    try:
        chunks_data = []
        if ext == "pdf":
            chunks_data = extract_pdf(content, doc_id, course_id)
        elif ext in ["ppt", "pptx"]:
            chunks_data = extract_pptx(file_path, doc_id, course_id)

        # Update progress
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if doc:
            doc.processing_progress = 50.0
            db.commit()

        # Create chunks and embeddings
        for i, cd in enumerate(chunks_data):
            chunk_id = f"chunk-{uuid.uuid4().hex[:8]}"
            chunk = DocumentChunk(
                id=chunk_id,
                document_id=doc_id,
                course_id=course_id,
                content=cd["content"],
                source_type=cd.get("source_type", ext),
                source_name=cd.get("source_name", "Uploaded Document"),
                page_number=cd.get("page_number"),
                slide_number=cd.get("slide_number"),
                topic=cd.get("topic", "General"),
                chunk_index=i
            )
            db.add(chunk)

            embedding = await rag_service.get_embedding(cd["content"])
            rag_service.index_chunk(chunk_id, {**cd, "course_id": course_id}, embedding)

        if doc:
            doc.status = "completed"
            doc.processing_progress = 100.0
            db.commit()

    except Exception as e:
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if doc:
            doc.status = "failed"
            db.commit()
        print(f"Error processing document: {e}")
    finally:
        db.close()


def extract_pdf(content: bytes, doc_id: str, course_id: str) -> list:
    """Extract text from PDF."""
    chunks = []
    try:
        import io
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(content))
        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text and text.strip():
                chunks.append({
                    "content": text.strip(),
                    "source_type": "pdf",
                    "source_name": "Uploaded PDF",
                    "page_number": page_num,
                    "topic": "General"
                })
    except Exception as e:
        print(f"PDF extraction error: {e}")
    return chunks


def extract_pptx(file_path: str, doc_id: str, course_id: str) -> list:
    """Extract text from PPTX."""
    chunks = []
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        for slide_num, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
            if texts:
                content = "\n".join(texts)
                chunks.append({
                    "content": content.strip(),
                    "source_type": "pptx",
                    "source_name": "Uploaded Slides",
                    "slide_number": slide_num,
                    "topic": "General"
                })
    except Exception as e:
        print(f"PPTX extraction error: {e}")
    return chunks


# ============================================================
# TOPICS
# ============================================================

@app.get("/api/topics/{course_id}")
async def get_topics(course_id: str, db: Session = Depends(get_db)):
    topics = db.query(Topic).filter(
        Topic.course_id == course_id,
        Topic.parent_topic_id == None
    ).order_by(Topic.order_index).all()

    result = []
    for t in topics:
        subtopics = db.query(Topic).filter(Topic.parent_topic_id == t.id).all()
        concepts = db.query(Concept).filter(Concept.topic_id == t.id).all()

        # Get sources for this topic
        chunks = db.query(DocumentChunk).filter(
            DocumentChunk.course_id == course_id,
            DocumentChunk.topic == t.name
        ).all()

        sources = []
        seen = set()
        for c in chunks:
            key = f"{c.source_type}-{c.source_name}-{c.page_number}-{c.slide_number}-{c.video_timestamp_start}"
            if key not in seen:
                seen.add(key)
                sources.append({
                    "source_type": c.source_type,
                    "source_name": c.source_name,
                    "page_number": c.page_number,
                    "slide_number": c.slide_number,
                    "timestamp_start": c.video_timestamp_start,
                    "timestamp_end": c.video_timestamp_end,
                    "content_preview": c.content[:100] + "..." if len(c.content) > 100 else c.content
                })

        # Get prerequisites
        prereqs = db.query(TopicPrerequisite).filter(
            TopicPrerequisite.topic_id == t.id
        ).all()
        prereq_names = []
        for p in prereqs:
            prereq_topic = db.query(Topic).filter(Topic.id == p.prerequisite_topic_id).first()
            if prereq_topic:
                prereq_names.append(prereq_topic.name)

        result.append({
            "id": t.id,
            "name": t.name,
            "unit": t.unit,
            "description": t.description,
            "difficulty": t.difficulty,
            "subtopics": [{"id": s.id, "name": s.name} for s in subtopics],
            "concepts": [{"id": c.id, "name": c.name} for c in concepts],
            "sources": sources,
            "prerequisites": prereq_names,
            "chunk_count": len(chunks)
        })

    return result


@app.get("/api/topics/{course_id}/graph")
async def get_topic_graph(course_id: str, db: Session = Depends(get_db)):
    """Get topic graph data for visualization."""
    topics = db.query(Topic).filter(
        Topic.course_id == course_id,
        Topic.parent_topic_id == None
    ).order_by(Topic.order_index).all()

    nodes = []
    edges = []

    mastery_records = db.query(Mastery).filter(
        Mastery.user_id == DEMO_USER_ID,
        Mastery.course_id == course_id
    ).all()
    mastery_map = {m.topic: m.mastery_score for m in mastery_records}

    for i, t in enumerate(topics):
        mastery = mastery_map.get(t.name, 0.5)
        nodes.append({
            "id": t.id,
            "label": t.name,
            "mastery": mastery,
            "unit": t.unit,
            "position": {"x": 300, "y": 100 + i * 150}
        })

    # Get prerequisites as edges
    prereqs = db.query(TopicPrerequisite).filter(
        TopicPrerequisite.topic_id.in_([t.id for t in topics])
    ).all()
    for p in prereqs:
        edges.append({
            "source": p.prerequisite_topic_id,
            "target": p.topic_id
        })

    return {"nodes": nodes, "edges": edges}


# ============================================================
# AI TUTOR
# ============================================================

@app.post("/api/tutor/query")
async def tutor_query(data: TutorQuery, db: Session = Depends(get_db)):
    """Handle a tutoring question with RAG."""
    # Retrieve relevant chunks
    chunks = await rag_service.retrieve(data.question, data.course_id, top_k=5)

    # Generate grounded answer
    response = await rag_service.generate_answer(data.question, chunks, data.language)

    # Save to session
    session_id = data.session_id
    if not session_id:
        session_id = f"session-{uuid.uuid4().hex[:8]}"
        session = TutorSession(
            id=session_id,
            user_id=DEMO_USER_ID,
            course_id=data.course_id
        )
        db.add(session)

    # Save user message
    user_msg = TutorMessage(
        id=f"msg-{uuid.uuid4().hex[:8]}",
        session_id=session_id,
        role="user",
        content=data.question
    )
    db.add(user_msg)

    # Save assistant message
    asst_msg = TutorMessage(
        id=f"msg-{uuid.uuid4().hex[:8]}",
        session_id=session_id,
        role="assistant",
        content=response["answer"],
        citations=response.get("citations"),
        is_grounded=response.get("is_grounded", True)
    )
    db.add(asst_msg)
    db.commit()

    return {
        "answer": response["answer"],
        "is_grounded": response.get("is_grounded", True),
        "citations": response.get("citations", []),
        "suggest_outside": response.get("suggest_outside", False),
        "session_id": session_id
    }


@app.get("/api/tutor/sessions")
async def get_tutor_sessions(course_id: str, db: Session = Depends(get_db)):
    sessions = db.query(TutorSession).filter(
        TutorSession.user_id == DEMO_USER_ID,
        TutorSession.course_id == course_id
    ).all()
    return [{
        "id": s.id,
        "created_at": s.created_at.isoformat() if s.created_at else None,
        "message_count": len(s.messages)
    } for s in sessions]


@app.get("/api/tutor/sessions/{session_id}/messages")
async def get_session_messages(session_id: str, db: Session = Depends(get_db)):
    messages = db.query(TutorMessage).filter(
        TutorMessage.session_id == session_id
    ).order_by(TutorMessage.created_at).all()
    return [{
        "id": m.id,
        "role": m.role,
        "content": m.content,
        "citations": m.citations,
        "is_grounded": m.is_grounded,
        "created_at": m.created_at.isoformat() if m.created_at else None
    } for m in messages]


# ============================================================
# ASSESSMENTS
# ============================================================

@app.post("/api/assessments/generate")
async def generate_assessment(data: AssessmentCreate, db: Session = Depends(get_db)):
    """Generate a new assessment with verified questions."""
    # Get relevant chunks
    all_chunks = []
    for chunk_id, chunk_data in rag_service.chunks.items():
        if chunk_data.get("course_id") == data.course_id:
            if not data.topics or chunk_data.get("topic") in data.topics:
                all_chunks.append(chunk_data)

    # Generate questions
    questions = await assessment_service.generate_questions(
        all_chunks, data.topics, data.difficulty,
        data.question_count, data.question_types
    )

    # Verify questions
    verified = []
    for q in questions:
        result = assessment_service.verify_question(q, all_chunks)
        if result["passed"]:
            verified.append(q)

    # If too few passed, include all
    if len(verified) < data.question_count // 2:
        verified = questions[:data.question_count]

    # Create assessment record
    assessment_id = f"assessment-{uuid.uuid4().hex[:8]}"
    assessment = Assessment(
        id=assessment_id,
        user_id=DEMO_USER_ID,
        course_id=data.course_id,
        title=f"{', '.join(data.topics)} Quiz",
        topics=data.topics,
        difficulty=data.difficulty,
        question_count=len(verified)
    )
    db.add(assessment)

    # Save questions
    for i, q in enumerate(verified):
        question = Question(
            id=q.get("id", f"q-{uuid.uuid4().hex[:8]}"),
            assessment_id=assessment_id,
            course_id=data.course_id,
            question_text=q["question_text"],
            question_type=q.get("question_type", "mcq"),
            options=q.get("options"),
            correct_answer=q.get("correct_answer", ""),
            explanation=q.get("explanation", ""),
            topic=q.get("topic", ""),
            subtopic=q.get("subtopic"),
            difficulty=q.get("difficulty", data.difficulty),
            source_citation=q.get("source_citation", ""),
            source_type=q.get("source_type"),
            source_location=q.get("source_location"),
            order_index=i
        )
        db.add(question)

    db.commit()

    return {
        "id": assessment_id,
        "title": assessment.title,
        "question_count": len(verified),
        "questions": [{
            "id": q.get("id"),
            "question_text": q["question_text"],
            "question_type": q.get("question_type", "mcq"),
            "options": q.get("options"),
            "topic": q.get("topic"),
            "difficulty": q.get("difficulty"),
            "order_index": i
        } for i, q in enumerate(verified)],
        "new_questions": len(verified),
        "repeated": 0
    }


@app.post("/api/assessments/{assessment_id}/submit")
async def submit_assessment(assessment_id: str, data: AnswerSubmit, db: Session = Depends(get_db)):
    """Submit and grade an assessment."""
    assessment = db.query(Assessment).filter(Assessment.id == assessment_id).first()
    if not assessment:
        raise HTTPException(404, "Assessment not found")

    questions = db.query(Question).filter(
        Question.assessment_id == assessment_id
    ).order_by(Question.order_index).all()

    results = []
    correct_count = 0
    topic_results = {}

    for q in questions:
        user_answer = data.answers.get(q.id, "")
        score_result = assessment_service.score_answer({
            "correct_answer": q.correct_answer,
            "question_type": q.question_type,
            "explanation": q.explanation,
            "source_citation": q.source_citation,
            "topic": q.topic
        }, user_answer)

        is_correct = score_result["is_correct"]
        if is_correct:
            correct_count += 1

        # Save attempt
        attempt = QuestionAttempt(
            id=f"attempt-{uuid.uuid4().hex[:8]}",
            question_id=q.id,
            user_id=DEMO_USER_ID,
            user_answer=user_answer,
            is_correct=is_correct
        )
        db.add(attempt)

        # Track topic performance
        if q.topic not in topic_results:
            topic_results[q.topic] = {"correct": 0, "total": 0}
        topic_results[q.topic]["total"] += 1
        if is_correct:
            topic_results[q.topic]["correct"] += 1

        # Update mastery
        mastery_record = db.query(Mastery).filter(
            Mastery.user_id == DEMO_USER_ID,
            Mastery.course_id == assessment.course_id,
            Mastery.topic == q.topic
        ).first()

        if mastery_record:
            update = learner_model_service.update_mastery(
                mastery_record.mastery_score,
                is_correct,
                q.difficulty or "medium",
                mastery_record.confidence
            )
            mastery_record.mastery_score = update["mastery_score"]
            mastery_record.confidence = update["confidence"]
            mastery_record.attempts += 1
            if is_correct:
                mastery_record.correct += 1
            mastery_record.last_reviewed = datetime.now(timezone.utc)
        else:
            initial = learner_model_service.update_mastery(0.5, is_correct, q.difficulty or "medium")
            mastery_record = Mastery(
                id=f"mastery-{uuid.uuid4().hex[:8]}",
                user_id=DEMO_USER_ID,
                course_id=assessment.course_id,
                topic=q.topic,
                mastery_score=initial["mastery_score"],
                confidence=initial["confidence"],
                attempts=1,
                correct=1 if is_correct else 0,
                last_reviewed=datetime.now(timezone.utc)
            )
            db.add(mastery_record)

        results.append({
            "question_id": q.id,
            "question_text": q.question_text,
            "question_type": q.question_type,
            "options": q.options,
            "user_answer": user_answer,
            "correct_answer": q.correct_answer,
            "is_correct": is_correct,
            "explanation": q.explanation,
            "source_citation": q.source_citation,
            "source_type": q.source_type,
            "source_location": q.source_location,
            "topic": q.topic,
            "difficulty": q.difficulty
        })

    # Update assessment
    score = (correct_count / len(questions) * 100) if questions else 0
    assessment.score = score
    assessment.completed = True
    assessment.completed_at = datetime.now(timezone.utc)
    db.commit()

    # Detect misconceptions
    all_attempts = [{
        "topic": r["topic"],
        "subtopic": r.get("source_citation", ""),
        "is_correct": r["is_correct"]
    } for r in results]
    misconceptions = learner_model_service.detect_misconceptions(all_attempts)

    # Topic performance
    topic_performance = []
    for topic, stats in topic_results.items():
        pct = (stats["correct"] / stats["total"] * 100) if stats["total"] > 0 else 0
        topic_performance.append({
            "topic": topic,
            "correct": stats["correct"],
            "total": stats["total"],
            "percentage": round(pct, 1)
        })

    return {
        "assessment_id": assessment_id,
        "score": round(score, 1),
        "correct": correct_count,
        "total": len(questions),
        "results": results,
        "topic_performance": topic_performance,
        "misconceptions": misconceptions
    }


@app.get("/api/assessments/{course_id}")
async def get_assessments(course_id: str, db: Session = Depends(get_db)):
    assessments = db.query(Assessment).filter(
        Assessment.user_id == DEMO_USER_ID,
        Assessment.course_id == course_id
    ).order_by(Assessment.created_at.desc()).all()

    return [{
        "id": a.id,
        "title": a.title,
        "topics": a.topics,
        "difficulty": a.difficulty,
        "question_count": a.question_count,
        "score": a.score,
        "completed": a.completed,
        "created_at": a.created_at.isoformat() if a.created_at else None,
        "completed_at": a.completed_at.isoformat() if a.completed_at else None
    } for a in assessments]


@app.get("/api/assessments/detail/{assessment_id}")
async def get_assessment_detail(assessment_id: str, db: Session = Depends(get_db)):
    assessment = db.query(Assessment).filter(Assessment.id == assessment_id).first()
    if not assessment:
        raise HTTPException(404, "Assessment not found")

    questions = db.query(Question).filter(
        Question.assessment_id == assessment_id
    ).order_by(Question.order_index).all()

    return {
        "id": assessment.id,
        "title": assessment.title,
        "topics": assessment.topics,
        "difficulty": assessment.difficulty,
        "score": assessment.score,
        "completed": assessment.completed,
        "questions": [{
            "id": q.id,
            "question_text": q.question_text,
            "question_type": q.question_type,
            "options": q.options,
            "correct_answer": q.correct_answer if assessment.completed else None,
            "explanation": q.explanation if assessment.completed else None,
            "source_citation": q.source_citation,
            "topic": q.topic,
            "difficulty": q.difficulty,
            "order_index": q.order_index
        } for q in questions]
    }


# ============================================================
# MASTERY
# ============================================================

@app.get("/api/mastery/{course_id}")
async def get_mastery(course_id: str, db: Session = Depends(get_db)):
    records = db.query(Mastery).filter(
        Mastery.user_id == DEMO_USER_ID,
        Mastery.course_id == course_id
    ).all()

    # Apply forgetting curves
    now = datetime.now(timezone.utc)
    result = []
    for m in records:
        adjusted = learner_model_service.apply_forgetting(
            m.mastery_score, m.last_reviewed, now
        )
        result.append({
            "topic": m.topic,
            "mastery_score": adjusted,
            "raw_mastery": m.mastery_score,
            "confidence": m.confidence,
            "attempts": m.attempts,
            "correct": m.correct,
            "last_reviewed": m.last_reviewed.isoformat() if m.last_reviewed else None
        })

    return result


@app.post("/api/mastery/update")
async def update_mastery(data: MasteryUpdate, db: Session = Depends(get_db)):
    record = db.query(Mastery).filter(
        Mastery.user_id == DEMO_USER_ID,
        Mastery.course_id == data.course_id,
        Mastery.topic == data.topic
    ).first()

    if not record:
        initial = learner_model_service.update_mastery(0.5, data.is_correct, data.difficulty)
        record = Mastery(
            id=f"mastery-{uuid.uuid4().hex[:8]}",
            user_id=DEMO_USER_ID,
            course_id=data.course_id,
            topic=data.topic,
            mastery_score=initial["mastery_score"],
            confidence=initial["confidence"],
            attempts=1,
            correct=1 if data.is_correct else 0,
            last_reviewed=datetime.now(timezone.utc)
        )
        db.add(record)
    else:
        update = learner_model_service.update_mastery(
            record.mastery_score, data.is_correct, data.difficulty, record.confidence
        )
        record.mastery_score = update["mastery_score"]
        record.confidence = update["confidence"]
        record.attempts += 1
        if data.is_correct:
            record.correct += 1
        record.last_reviewed = datetime.now(timezone.utc)

    db.commit()
    return {"topic": data.topic, "mastery_score": record.mastery_score, "confidence": record.confidence}


# ============================================================
# STUDY PLAN
# ============================================================

@app.get("/api/study-plan/{course_id}")
async def get_study_plan(course_id: str, exam_date: Optional[str] = None,
                          daily_hours: float = 2.0, db: Session = Depends(get_db)):
    mastery_records = db.query(Mastery).filter(
        Mastery.user_id == DEMO_USER_ID,
        Mastery.course_id == course_id
    ).all()

    mastery_data = [{
        "topic": m.topic,
        "mastery_score": m.mastery_score,
        "confidence": m.confidence,
        "last_reviewed": m.last_reviewed.isoformat() if m.last_reviewed else None
    } for m in mastery_records]

    # Get prerequisites
    topics = db.query(Topic).filter(
        Topic.course_id == course_id,
        Topic.parent_topic_id == None
    ).all()
    prereqs_map = {}
    for t in topics:
        tp = db.query(TopicPrerequisite).filter(TopicPrerequisite.topic_id == t.id).all()
        prereq_names = []
        for p in tp:
            pt = db.query(Topic).filter(Topic.id == p.prerequisite_topic_id).first()
            if pt:
                prereq_names.append(pt.name)
        prereqs_map[t.name] = prereq_names

    priorities = learner_model_service.generate_study_priorities(mastery_data, prereqs_map)

    # Generate daily plan
    daily_minutes = int(daily_hours * 60)
    today_plan = []
    remaining = daily_minutes
    for p in priorities:
        if remaining <= 0:
            break
        mins = min(p["time_minutes"], remaining)
        today_plan.append({
            "topic": p["topic"],
            "action": p["action"],
            "time_minutes": mins,
            "reason": p["reason"],
            "mastery": p["mastery"]
        })
        remaining -= mins

    # Days until exam
    days_left = None
    if exam_date:
        try:
            exam = datetime.fromisoformat(exam_date)
            days_left = (exam - datetime.now()).days
        except (ValueError, TypeError):
            pass

    return {
        "daily_plan": today_plan,
        "priorities": priorities,
        "total_study_time": daily_minutes - remaining,
        "exam_date": exam_date,
        "days_until_exam": days_left,
        "weak_topics": [p for p in priorities if p["mastery"] < 0.5],
        "strong_topics": [p for p in priorities if p["mastery"] >= 0.75]
    }


@app.post("/api/study-plan/{course_id}")
async def save_study_plan(course_id: str, data: StudyPlanCreate, db: Session = Depends(get_db)):
    plan = StudyPlan(
        id=f"plan-{uuid.uuid4().hex[:8]}",
        user_id=DEMO_USER_ID,
        course_id=course_id,
        exam_date=data.exam_date,
        daily_hours=data.daily_hours,
        plan_data={}
    )
    db.add(plan)
    db.commit()
    return {"id": plan.id, "status": "saved"}


# ============================================================
# REVISION MATERIAL
# ============================================================

@app.post("/api/revision/generate")
async def generate_revision(data: RevisionRequest, db: Session = Depends(get_db)):
    """Generate revision material from course content."""
    # Get chunks for topic
    chunks = []
    for chunk_id, chunk_data in rag_service.chunks.items():
        if chunk_data.get("course_id") == data.course_id and chunk_data.get("topic") == data.topic:
            chunks.append(chunk_data)

    if not chunks:
        raise HTTPException(404, "No content found for this topic")

    if data.type == "flashcards":
        cards = []
        for c in chunks[:8]:
            content = c.get("content", "")
            # Create Q&A flashcard
            concept = c.get("concept", c.get("subtopic", "Concept"))
            source = f"{c.get('source_name', 'Source')}"
            if c.get("page_number"):
                source += f" — Page {c['page_number']}"
            elif c.get("slide_number"):
                source += f" — Slide {c['slide_number']}"

            # Generate front and back
            sentences = content.split(".")
            if len(sentences) >= 2:
                front = f"What is {concept}?"
                back = sentences[0].strip() + "."
            else:
                front = f"Explain {concept}"
                back = content[:200]

            cards.append({
                "front": front,
                "back": back,
                "source": source,
                "topic": data.topic,
                "concept": concept
            })
        return {"type": "flashcards", "topic": data.topic, "cards": cards}

    elif data.type == "notes":
        notes = f"# Quick Notes: {data.topic}\n\n"
        for c in chunks:
            concept = c.get("concept", c.get("subtopic", ""))
            source = f"{c.get('source_name', 'Source')}"
            if c.get("page_number"):
                source += f" — Page {c['page_number']}"
            elif c.get("slide_number"):
                source += f" — Slide {c['slide_number']}"
            notes += f"## {concept}\n{c['content']}\n*Source: {source}*\n\n"
        return {"type": "notes", "topic": data.topic, "content": notes}

    elif data.type == "summary":
        summary_parts = []
        for c in chunks[:5]:
            content = c.get("content", "")
            # Take first 2 sentences as summary
            sentences = content.split(".")
            summary = ". ".join(s.strip() for s in sentences[:2] if s.strip()) + "."
            summary_parts.append(summary)
        return {
            "type": "summary",
            "topic": data.topic,
            "content": "\n\n".join(summary_parts),
            "source_count": len(chunks)
        }

    raise HTTPException(400, f"Unknown revision type: {data.type}")


# ============================================================
# SOURCE VIEWER
# ============================================================

@app.get("/api/sources/{course_id}")
async def get_sources(course_id: str, source_type: Optional[str] = None,
                       topic: Optional[str] = None, db: Session = Depends(get_db)):
    """Get source content for viewing."""
    query = db.query(DocumentChunk).filter(DocumentChunk.course_id == course_id)
    if source_type:
        query = query.filter(DocumentChunk.source_type == source_type)
    if topic:
        query = query.filter(DocumentChunk.topic == topic)

    chunks = query.all()
    return [{
        "id": c.id,
        "content": c.content,
        "source_type": c.source_type,
        "source_name": c.source_name,
        "page_number": c.page_number,
        "slide_number": c.slide_number,
        "timestamp_start": c.video_timestamp_start,
        "timestamp_end": c.video_timestamp_end,
        "topic": c.topic,
        "subtopic": c.subtopic,
        "concept": c.concept
    } for c in chunks]


# ============================================================
# EVALUATION
# ============================================================

@app.post("/api/evaluation/run")
async def run_evaluation(course_id: str = DEMO_COURSE_ID, db: Session = Depends(get_db)):
    """Run RAG evaluation pipeline."""
    results = await evaluation_service.run_evaluation(
        EVALUATION_QUESTIONS, rag_service, course_id
    )

    # Get mastery for personalization eval
    mastery_records = db.query(Mastery).filter(
        Mastery.user_id == DEMO_USER_ID,
        Mastery.course_id == course_id
    ).all()
    mastery_data = [{
        "topic": m.topic,
        "mastery_score": m.mastery_score,
    } for m in mastery_records]

    personalization = evaluation_service.evaluate_personalization(mastery_data)

    # Save result
    eval_result = EvaluationResult(
        id=f"eval-{uuid.uuid4().hex[:8]}",
        course_id=course_id,
        faithfulness=results["faithfulness"],
        answer_relevancy=results["answer_relevancy"],
        context_precision=results["context_precision"],
        context_recall=results["context_recall"],
        mastery_gain=personalization["mastery_gain"],
        question_repetition_rate=personalization["question_repetition_rate"],
        weak_topic_improvement=personalization["weak_topic_improvement"],
        details=results["details"]
    )
    db.add(eval_result)
    db.commit()

    return {
        **results,
        "personalization": personalization
    }


@app.get("/api/evaluation/results")
async def get_evaluation_results(course_id: str = DEMO_COURSE_ID, db: Session = Depends(get_db)):
    results = db.query(EvaluationResult).filter(
        EvaluationResult.course_id == course_id
    ).order_by(EvaluationResult.created_at.desc()).all()

    if not results:
        return None

    latest = results[0]
    return {
        "faithfulness": latest.faithfulness,
        "answer_relevancy": latest.answer_relevancy,
        "context_precision": latest.context_precision,
        "context_recall": latest.context_recall,
        "mastery_gain": latest.mastery_gain,
        "question_repetition_rate": latest.question_repetition_rate,
        "weak_topic_improvement": latest.weak_topic_improvement,
        "details": latest.details,
        "created_at": latest.created_at.isoformat() if latest.created_at else None,
        "total_evaluations": len(results)
    }


# ============================================================
# DASHBOARD
# ============================================================

@app.get("/api/dashboard")
async def get_dashboard(course_id: str = DEMO_COURSE_ID, db: Session = Depends(get_db)):
    """Get dashboard data."""
    # Mastery
    mastery_records = db.query(Mastery).filter(
        Mastery.user_id == DEMO_USER_ID,
        Mastery.course_id == course_id
    ).all()

    overall_mastery = 0
    topics_mastered = 0
    topics_weak = []
    topic_data = []
    total_attempts = 0
    total_correct = 0

    for m in mastery_records:
        adjusted = learner_model_service.apply_forgetting(
            m.mastery_score, m.last_reviewed
        )
        if adjusted >= 0.75:
            topics_mastered += 1
        if adjusted < 0.6:
            topics_weak.append({"topic": m.topic, "mastery": round(adjusted * 100, 1)})
        topic_data.append({"topic": m.topic, "mastery": round(adjusted * 100, 1)})
        total_attempts += m.attempts
        total_correct += m.correct
        overall_mastery += adjusted

    if mastery_records:
        overall_mastery /= len(mastery_records)

    # Recent assessments
    recent = db.query(Assessment).filter(
        Assessment.user_id == DEMO_USER_ID,
        Assessment.course_id == course_id,
        Assessment.completed == True
    ).order_by(Assessment.completed_at.desc()).limit(5).all()

    recent_assessments = [{
        "id": a.id,
        "title": a.title,
        "score": a.score,
        "topics": a.topics,
        "difficulty": a.difficulty,
        "date": a.completed_at.isoformat() if a.completed_at else None
    } for a in recent]

    # Recommendations
    recommendations = []
    topics_weak_sorted = sorted(topics_weak, key=lambda x: x["mastery"])
    for tw in topics_weak_sorted[:3]:
        recommendations.append({
            "action": f"Revise {tw['topic']}",
            "reason": f"Your mastery is at {tw['mastery']}%. Focus on building stronger foundations.",
            "topic": tw["topic"],
            "priority": "high" if tw["mastery"] < 50 else "medium"
        })

    return {
        "user": {"name": "Thanusha", "id": DEMO_USER_ID},
        "overall_mastery": round(overall_mastery * 100, 1),
        "topics_mastered": topics_mastered,
        "topics_total": len(mastery_records),
        "topics_needing_attention": len(topics_weak),
        "questions_answered": total_attempts,
        "study_streak": 5,  # Demo value
        "weak_areas": topics_weak_sorted,
        "topic_mastery": topic_data,
        "recent_assessments": recent_assessments,
        "recommendations": recommendations
    }


# ============================================================
# HEALTH
# ============================================================

@app.get("/api/health")
async def health():
    return {
        "status": "ok",
        "has_openai": rag_service.has_api,
        "demo_mode": not rag_service.has_api,
        "chunks_indexed": len(rag_service.chunks)
    }


import random
